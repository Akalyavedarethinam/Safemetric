"""Editable Word (.docx) inspection report."""
from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from ..models import Inspection
from .common import DISCLAIMER, FIELD_STATUS_LABEL, LEGEND, STATUS_HEX, STATUS_LABEL, image_bytes, meta


def _shade(cell, hexcolor: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor.lstrip("#"))
    tcPr.append(shd)


def _rgb(hexcolor):
    h = hexcolor.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(8.5), RGBColor(255, 255, 255)
        _shade(c, "#0b1f3a")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            color = None
            if isinstance(v, tuple):
                v, color = v
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(8)
            if color:
                run.bold, run.font.color.rgb = True, _rgb(color)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Mm(w)
    return t


def build(insp: Inspection, generated_by: str = "") -> bytes:
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Mm(15)
        s.top_margin = s.bottom_margin = Mm(14)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(9.5)

    h = doc.add_heading("SAFEMETRIC", 0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LEGAL METROLOGY COMPLIANCE INSPECTION REPORT\n")
    r.bold, r.font.color.rgb = True, _rgb("#1d4ed8")
    p.add_run("The Legal Metrology Act, 2009 · Legal Metrology (Packaged Commodities) Rules, 2011").font.size = Pt(8)

    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in meta(insp):
        cells = t.add_row().cells
        cells[0].text = k
        cells[0].paragraphs[0].runs[0].bold = True
        _shade(cells[0], "#f1f5f9")
        cells[1].text = str(v)
        if k == "Overall status":
            run = cells[1].paragraphs[0].runs[0]
            run.bold, run.font.color.rgb = True, RGBColor(255, 255, 255)
            _shade(cells[1], STATUS_HEX.get(insp.overall_status, "#334155"))
    doc.add_paragraph(f"Checks - passed: {insp.passed} · non-compliant: {insp.failed} · needs verification: "
                      f"{insp.needs_verification} · not applicable: {insp.not_applicable}")
    if (insp.summary or {}).get("overall_reason"):
        doc.add_paragraph("Basis of overall status: " + insp.summary["overall_reason"])

    doc.add_heading("1. Package photographs", level=1)
    for img in insp.images:
        data = image_bytes(img, max_side=800, boxes=[f.bbox for f in insp.fields if f.source_image_id == img.id and f.bbox])
        if data:
            doc.add_picture(io.BytesIO(data), width=Mm(70))
        doc.add_paragraph(f"{img.side.upper()} - {img.original_filename} · {img.width}×{img.height}px · quality "
                          f"{img.quality_status} · OCR {img.ocr.mean_confidence if img.ocr else 0:.1f}% · SHA-256 {img.sha256}")

    doc.add_heading("2. Declarations extracted (OCR - detected, not verified)", level=1)
    _table(doc, ["Declaration", "Extracted value", "Confidence", "Source", "Status"],
           [[f.label, f.value or "- not detected -", f"{f.confidence:.0f}%", f"{f.source_side} {f.source_image}".strip() or "-",
             (FIELD_STATUS_LABEL.get(f.status, f.status), STATUS_HEX.get(f.status, "#334155"))] for f in insp.fields],
           [40, 60, 18, 35, 27])

    doc.add_heading("3. Rule findings", level=1)
    _table(doc, ["Rule", "Check", "Severity", "Result", "Reason"],
           [[f.rule_no, f.title, f.severity, (STATUS_LABEL.get(f.status, f.status), STATUS_HEX.get(f.status, "#334155")),
             f.reason] for f in insp.findings], [12, 45, 16, 25, 82])

    doc.add_heading("4. Required actions", level=1)
    recs = insp.recommendations or []
    if not recs:
        doc.add_paragraph("No actions required.")
    for rec in recs:
        doc.add_paragraph(f"[{rec['type'].replace('_', ' ')} · {rec['urgency']}] {rec['title']} - {rec['action']} "
                          f"({rec['reference']})", style="List Bullet")

    doc.add_heading("5. OCR text", level=1)
    for img in insp.images:
        doc.add_paragraph(f"{img.side.upper()} - {img.original_filename}").runs[0].bold = True
        doc.add_paragraph(img.ocr.full_text if img.ocr else "(no OCR result)")

    doc.add_heading("Officer remarks", level=1)
    doc.add_paragraph(insp.notes or "\n\n")
    doc.add_heading("Legend", level=2)
    for k, v in LEGEND:
        doc.add_paragraph(f"{k}: {v}", style="List Bullet")
    doc.add_paragraph(DISCLAIMER).runs[0].font.size = Pt(8)
    doc.add_paragraph(f"Integrity hash: {insp.integrity_hash}").runs[0].font.size = Pt(7)
    doc.add_paragraph(f"\n\nSignature of Inspecting Officer: ____________________     Name: {insp.inspector_name}     Date: __________")
    if generated_by:
        doc.add_paragraph(f"Generated by {generated_by}").runs[0].font.size = Pt(7)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
