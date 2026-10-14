"""Editable Excel (.xlsx) inspection workbook: Summary, Declarations, Findings, OCR."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import Inspection
from .common import FIELD_STATUS_LABEL, STATUS_HEX, STATUS_LABEL, meta

HEAD = PatternFill("solid", fgColor="0B1F3A")


def _sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), HEAD, Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build(insp: Inspection) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _sheet(ws, ["Item", "Value"], meta(insp) + [
        ("Passed checks", insp.passed), ("Non-compliant checks", insp.failed),
        ("Needs verification", insp.needs_verification), ("Not applicable", insp.not_applicable),
        ("Basis of overall status", (insp.summary or {}).get("overall_reason", "")),
        ("Integrity hash", insp.integrity_hash)], [32, 90])
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "Overall status":
            row[1].font = Font(bold=True, color="FFFFFF")
            row[1].fill = PatternFill("solid", fgColor=STATUS_HEX.get(insp.overall_status, "#334155").lstrip("#"))

    ws = wb.create_sheet("Declarations")
    _sheet(ws, ["Field", "Declaration", "Extracted value", "Confidence %", "Status", "Source side", "Source image",
                "OCR evidence", "BBox x", "BBox y", "BBox w", "BBox h"],
           [[f.field, f.label, f.value, f.confidence, FIELD_STATUS_LABEL.get(f.status, f.status), f.source_side,
             f.source_image, f.evidence] + ([f.bbox["x"], f.bbox["y"], f.bbox["width"], f.bbox["height"]] if f.bbox else [None] * 4)
            for f in insp.fields], [18, 30, 45, 12, 18, 12, 22, 60, 8, 8, 8, 8])

    ws = wb.create_sheet("Rule findings")
    _sheet(ws, ["Check ID", "Rule", "Check", "Severity", "Result", "Reason", "Recommendation", "Verification",
                "Applicability", "Condition", "Reference"],
           [[f.check_id, f.rule_no, f.title, f.severity, STATUS_LABEL.get(f.status, f.status), f.reason, f.recommendation,
             f.verification, f.applicability, f.condition, f.reference] for f in insp.findings],
           [20, 8, 40, 10, 18, 70, 45, 40, 30, 50, 40])
    for row in ws.iter_rows(min_row=2):
        hexc = {"Pass": "15803D", "Non-compliant": "B91C1C", "Needs verification": "B45309"}.get(row[4].value)
        if hexc:
            row[4].font = Font(bold=True, color=hexc)

    ws = wb.create_sheet("OCR text")
    rows = []
    for img in insp.images:
        for line in (img.ocr.lines if img.ocr else []):
            b = line["bbox"]
            rows.append([img.side, img.original_filename, line["text"], line["confidence"], b["x"], b["y"],
                         b["width"], b["height"], line.get("pass", "")])
    _sheet(ws, ["Side", "Image", "Text line", "Confidence %", "x", "y", "w", "h", "OCR pass"], rows,
           [10, 22, 70, 12, 7, 7, 7, 7, 16])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
